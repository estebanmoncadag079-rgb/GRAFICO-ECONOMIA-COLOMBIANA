"""Actualiza PIB trimestral desde los anexos de produccion del DANE.

El DANE revisa periodos anteriores. Por eso se sustituye toda la serie
verificada en cada publicacion; nunca se crean trimestres desde datos anuales.
"""

import argparse
import io
import re
from pathlib import Path
from urllib.parse import urljoin, urlparse

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
CSV_PIB = BASE_DIR / "pib_colombia.csv"
DANE_PAGE = (
    "https://www.dane.gov.co/index.php/estadisticas-por-tema/"
    "cuentas-nacionales/cuentas-nacionales-trimestrales/pib-informacion-tecnica"
)
QUARTERS = {"I": 1, "II": 4, "III": 7, "IV": 10}


def annex_links(html):
    soup = BeautifulSoup(html, "html.parser")
    result = {}
    for kind, marker in {
        "real": "anex-ProduccionConstantes-",
        "nominal": "anex-ProduccionCorriente-",
    }.items():
        links = [a for a in soup.select("a[href]") if marker.lower() in a["href"].lower()]
        if len(links) != 1:
            raise ValueError(f"Se esperaba 1 anexo {kind} vigente; encontrados: {len(links)}")
        anchor = links[0]
        url = urljoin(DANE_PAGE, anchor["href"])
        if urlparse(url).hostname != "www.dane.gov.co" or not url.lower().endswith(".xlsx"):
            raise ValueError(f"URL de anexo no valida: {url}")
        row = anchor.find_parent("tr")
        if row is None:
            raise ValueError(f"Falta fecha de publicacion para {kind}")
        match = pd.Series([row.get_text(" ", strip=True)]).str.extract(
            r"(\d{1,2}/\d{1,2}/\d{4})"
        )[0].iloc[0]
        if pd.isna(match):
            raise ValueError(f"Falta fecha de publicacion para {kind}")
        result[kind] = (url, pd.to_datetime(match, format="%d/%m/%Y"))
    if result["real"][1] != result["nominal"][1]:
        raise ValueError("Los anexos reales y nominales tienen fechas distintas")
    return result


def read_pib_levels(content, sheet_name):
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"No existe {sheet_name} en el anexo")
        rows = list(wb[sheet_name].values)
        if "producto interno bruto" not in str(rows[4][0]).lower():
            raise ValueError("El anexo no parece ser PIB")
        expected = "originales" if sheet_name == "Cuadro 1" else "ajustados"
        if expected not in str(rows[7][0]).lower():
            raise ValueError(f"{sheet_name} ya no contiene datos {expected}")
        candidates = [
            r for r in rows
            if str(r[0]).strip() == "B.1b"
            and str(r[2]).strip().lower() == "producto interno bruto"
            and sum(isinstance(v, (int, float)) for v in r[3:]) > 30
            and isinstance(r[3], (int, float)) and r[3] > 1000
        ]
        if len(candidates) != 1:
            raise ValueError(f"Fila de niveles PIB ambigua: {len(candidates)}")
        years, quarters, values = rows[11], rows[12], candidates[0]
        output = {}
        year = None
        for col in range(3, min(len(years), len(quarters), len(values))):
            year_match = re.match(r"^(20\d{2}|19\d{2})(?:p|pr)?$", str(years[col]).strip())
            if year_match:
                year = int(year_match.group(1))
            quarter = str(quarters[col]).strip()
            value = values[col]
            if quarter not in QUARTERS or not isinstance(value, (int, float)):
                continue
            if year is None or value <= 0:
                raise ValueError("Fecha o nivel PIB invalido")
            output[pd.Timestamp(year, QUARTERS[quarter], 1)] = float(value)
        if len(output) < 60:
            raise ValueError(f"Historia PIB demasiado corta: {len(output)} trimestres")
        return pd.Series(output, dtype="float64").sort_index()
    finally:
        wb.close()


def build_table(real_original, real_adjusted, nominal, publication, source_url):
    if not real_original.index.equals(real_adjusted.index) or not real_original.index.equals(nominal.index):
        raise ValueError("Los anexos PIB no cubren los mismos trimestres")
    expected = pd.date_range(real_original.index.min(), real_original.index.max(), freq="QS")
    if not real_original.index.equals(expected):
        raise ValueError("Faltan trimestres en el PIB")
    df = pd.DataFrame({
        "pib_real_miles_millones_ref2015": real_original,
        "pib_real_ajustado_miles_millones_ref2015": real_adjusted,
        "pib_nominal_billones_cop": nominal / 1000,
    })
    df["pib_real_yoy"] = real_original.pct_change(4, fill_method=None) * 100
    df["pib_real_qoq_sa"] = real_adjusted.pct_change(1, fill_method=None) * 100
    df["pib_nominal_yoy"] = nominal.pct_change(4, fill_method=None) * 100
    if df.index.max() > pd.Timestamp.today().normalize():
        raise ValueError("El anexo incluye trimestres futuros")
    if (df["pib_real_yoy"].abs() > 40).any() or (df["pib_real_qoq_sa"].abs() > 30).any():
        raise ValueError("Crecimiento PIB fuera de rango de control")
    df = df.loc["2009-01-01":].copy()
    df.insert(0, "fecha", df.index.strftime("%Y-%m-%d"))
    df.insert(1, "trimestre", [f"T{d.quarter} {d.year}" for d in df.index])
    df["fecha_publicacion_vintage"] = publication.strftime("%Y-%m-%d")
    df["fuente"] = source_url
    df["estado_dato"] = "DANE_publicado_sujeto_a_revision"
    return df.reset_index(drop=True)


def update(session=None):
    session = session or requests.Session()
    response = session.get(DANE_PAGE, timeout=40)
    response.raise_for_status()
    links = annex_links(response.text)
    content = {}
    for kind, (url, _) in links.items():
        response = session.get(url, timeout=90)
        response.raise_for_status()
        content[kind] = response.content
    real = read_pib_levels(content["real"], "Cuadro 1")
    adjusted = read_pib_levels(content["real"], "Cuadro 4")
    nominal = read_pib_levels(content["nominal"], "Cuadro 1")
    table = build_table(real, adjusted, nominal, links["real"][1], links["real"][0])
    name_match = re.search(r"-(I{1,3}|IV)trim(20\d{2})\.xlsx$", links["real"][0], re.I)
    if name_match:
        expected_last = pd.Timestamp(int(name_match.group(2)), QUARTERS[name_match.group(1).upper()], 1)
        if pd.Timestamp(table.iloc[-1]["fecha"]) != expected_last:
            raise ValueError("El ultimo trimestre del anexo no coincide con su nombre")
    if CSV_PIB.exists():
        old = pd.read_csv(CSV_PIB)
        new = pd.read_csv(io.StringIO(table.to_csv(index=False, float_format="%.8f")))
        if old.equals(new):
            print("PIB: sin cambios")
            return table
    table.to_csv(CSV_PIB, index=False, float_format="%.8f")
    print(f"PIB DANE: {len(table)} trimestres hasta {table.iloc[-1]['trimestre']}; publicado {links['real'][1].date()}")
    return table


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--estado", action="store_true", help="Mostrar tabla existente")
    args = parser.parse_args()
    if args.estado:
        print(pd.read_csv(CSV_PIB).tail(8).to_string(index=False))
    else:
        update()
